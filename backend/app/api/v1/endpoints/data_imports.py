import csv
import io
import json
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from sqlalchemy.orm import Session

from app.core.deps import get_current_active_user
from app.db.session import get_db
from app.models.capture import Capture
from app.models.data_import import DataImport
from app.models.site import SiteSentinelle
from app.models.user import User

router = APIRouter()

MAX_FILE_SIZE = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {".csv", ".json", ".xlsx"}


def _read_rows(filename: str, content: bytes) -> list[dict]:
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(400, "Le fichier CSV doit être encodé en UTF-8") from exc
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [dict(row) for row in csv.DictReader(io.StringIO(text), dialect=dialect)]

    if extension == ".json":
        try:
            payload = json.loads(content.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise HTTPException(400, "Le fichier JSON est invalide") from exc
        rows = payload.get("captures") if isinstance(payload, dict) else payload
        if not isinstance(rows, list) or not all(isinstance(row, dict) for row in rows):
            raise HTTPException(400, "Le JSON doit contenir une liste d'objets ou une clé 'captures'")
        return rows

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(503, "Le support Excel n'est pas installé sur le serveur") from exc
    try:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(values)]
        return [dict(zip(headers, row)) for row in values if any(value is not None for value in row)]
    except Exception as exc:
        raise HTTPException(400, "Le fichier Excel est illisible") from exc


def _parse_date(value) -> datetime:
    if isinstance(value, datetime):
        return value
    text = str(value or "").strip().replace("Z", "+00:00")
    if not text:
        raise ValueError("date_capture obligatoire")
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%Y %H:%M"):
            try:
                return datetime.strptime(text, pattern)
            except ValueError:
                pass
    raise ValueError("date_capture invalide")


@router.post("/upload")
async def upload_data(
    file: UploadFile,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    filename = Path(file.filename or "import").name
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, "Format non supporté. Utilisez CSV, JSON ou XLSX")
    content = await file.read(MAX_FILE_SIZE + 1)
    if not content:
        raise HTTPException(400, "Le fichier est vide")
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(413, "Le fichier dépasse la limite de 10 Mo")

    rows = _read_rows(filename, content)
    history = DataImport(
        filename=filename,
        file_type=extension.lstrip("."),
        statut="en_cours",
        total=len(rows),
        utilisateur_id=current_user.id,
    )
    db.add(history)
    db.flush()

    sites_by_code = {
        site.code.strip().lower(): site.id
        for site in db.query(SiteSentinelle).all()
        if site.code
    }
    errors = []
    captures = []
    for index, row in enumerate(rows, start=2):
        try:
            normalized = {str(key).strip().lower(): value for key, value in row.items() if key}
            site_id = normalized.get("site_id")
            code_site = str(normalized.get("code_site") or normalized.get("site") or "").strip().lower()
            if site_id not in (None, ""):
                site_id = int(site_id)
                if not db.get(SiteSentinelle, site_id):
                    raise ValueError(f"site_id {site_id} introuvable")
            else:
                site_id = sites_by_code.get(code_site)
                if not site_id:
                    raise ValueError(f"code_site '{code_site or '?'}' introuvable")
            espece = str(normalized.get("espece") or normalized.get("identification_ia") or "").strip()
            if not espece:
                raise ValueError("espece obligatoire")
            quantity = normalized.get("nombre_individus", normalized.get("quantite", 1))
            captures.append(Capture(
                site_id=site_id,
                date_capture=_parse_date(normalized.get("date_capture")),
                espece=espece,
                sexe=str(normalized.get("sexe") or "").strip() or None,
                stade=str(normalized.get("stade") or "").strip() or None,
                methode_capture=str(normalized.get("methode_capture") or "").strip() or None,
                nombre_individus=max(1, int(quantity or 1)),
                notes=str(normalized.get("notes") or "").strip() or None,
                utilisateur_id=current_user.id,
            ))
        except (TypeError, ValueError) as exc:
            errors.append(f"Ligne {index}: {exc}")

    db.add_all(captures)
    history.imported = len(captures)
    history.rejected = len(errors)
    history.errors_json = json.dumps(errors, ensure_ascii=False)
    history.statut = "succes" if not errors else ("partiel" if captures else "erreur")
    db.commit()
    db.refresh(history)
    return {
        "id": history.id,
        "message": f"{history.imported} capture(s) importée(s), {history.rejected} rejetée(s)",
        "total": history.total,
        "imported": history.imported,
        "rejected": history.rejected,
        "statut": history.statut,
        "errors": errors,
    }


@router.get("/history")
def import_history(
    limit: int = 30,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user),
):
    query = db.query(DataImport)
    if not current_user.is_superuser:
        query = query.filter(DataImport.utilisateur_id == current_user.id)
    imports = query.order_by(DataImport.created_at.desc()).limit(min(max(limit, 1), 100)).all()
    return [{
        "id": item.id,
        "filename": item.filename,
        "date": item.created_at,
        "total": item.total,
        "imported": item.imported,
        "rejected": item.rejected,
        "statut": item.statut,
        "errors": json.loads(item.errors_json or "[]"),
    } for item in imports]
